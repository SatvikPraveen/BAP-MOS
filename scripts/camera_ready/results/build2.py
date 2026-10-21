"""Add a prostate (TRUS) sheet and a bladder (PFUS1) sheet in front of the full workbook.

Each sheet stacks that dataset's tables (mean ± SD over seeds 42/43/44), its paired tests
and its notes. Built from collected.json with the same labels as build.py.
"""
import json
import math
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import labels as L  # noqa: E402

OUT = Path(sys.argv[1])
META = json.loads(Path(sys.argv[2]).read_text())
FULL = Path(sys.argv[3])  # workbook from build.py; the two sheets are added in front
d = L.d

TITLE = Font(bold=True, size=13)
HDR = Font(bold=True)
HFILL = PatternFill("solid", fgColor="DDE5F0")
PEND = Font(color="9C0006", italic=True)
NOTE = Font(italic=True, color="444444")


class Page:
    def __init__(self, ws):
        self.ws, self.r = ws, 1

    def text(self, t, font=None):
        c = self.ws.cell(row=self.r, column=1, value=t)
        if font:
            c.font = font
        self.r += 1

    def block(self, title, notes, rows, num_fmt=None):
        """Section: title, notes, header row (rows[0]) and data rows."""
        self.text(title, TITLE)
        for n in notes:
            self.text(n, NOTE)
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                c = self.ws.cell(row=self.r, column=j + 1, value=v)
                if i == 0:
                    c.font, c.fill = HDR, HFILL
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                elif v == "pending":
                    c.font = PEND
                elif num_fmt and j in num_fmt and isinstance(v, (int, float)):
                    c.number_format = num_fmt[j]
            self.r += 1
        self.r += 2


def hdr_cols(ds):
    return [f"{o if o != 'ALL' else 'All organs'} {L.MLABEL[m]}" + ("" if m == "dice" else f" ({L.UNIT[ds]})")
            for o in L.ORGANS[ds] for m in L.agg.METRICS]


def rows_models(ds, block, cond, models):
    rows = [["Model", ""] + hdr_cols(ds)]
    for m in models:
        rows.append([L.NAMES[m], ""] + [L.cell(ds, block, cond, m, o, me)
                                        for o in L.ORGANS[ds] for me in L.agg.METRICS])
    return rows


def rows_conditions(ds, models, conds=L.ROBUST):
    rows = [["Boxes", "Model"] + hdr_cols(ds)]
    for c in conds:
        for m in models:
            blk = "exact" if c == "exact" else "robust"
            rows.append([L.COND[c], L.NAMES[m]] + [L.cell(ds, blk, c, m, o, me)
                                                   for o in L.ORGANS[ds] for me in L.agg.METRICS])
    return rows


def rows_paired(ds):
    rows = [["Boxes", "A", "B", "Organ", "Metric", "n pairs", "A mean", "B mean", "Median diff (A-B)",
             "Fraction A better", "p (Wilcoxon)", "p (Holm)", "A invalid", "B invalid", "Status"]]
    nan = lambda x: None if x is None or (isinstance(x, float) and math.isnan(x)) else x
    for name, p in d["paired"].items():
        pds, cond, _ = name.split("|")
        if pds != ds:
            continue
        a, b = L.NAMES[p["a"].split("|")[3]], L.NAMES[p["b"].split("|")[3]]
        boxes = L.COND.get(cond.replace("loco_", ""), cond)
        if p["status"] != "ok":
            rows.append([boxes, a, b] + [None] * 11 + ["pending"])
            continue
        for r in p["rows"]:
            rows.append([boxes, a, b, "All organs" if r["organ"] == "ALL" else r["organ"], L.MLABEL[r["metric"]],
                         r["n_pairs"], nan(r["A_mean"]), nan(r["B_mean"]), nan(r["median_diff_a_minus_b"]),
                         nan(r["a_better_frac"]), nan(r["p_wilcoxon"]), nan(r["p_holm"]),
                         r["a_invalid"], r["b_invalid"], "ok"])
    return rows


PFMT = {6: "0.0000", 7: "0.0000", 8: "0.0000", 9: "0.000", 10: "0.00E+00", 11: "0.00E+00"}
COMMON = ("Every cell: mean ± sample SD (ddof=1) over seeds 42/43/44. 'pending' = not all 3 seeds finished "
          "(never a partial average). Per seed: mean over slice-organ pairs with a valid boundary.")
TRAIN = ("Train-time box jitter: every box-prompted model 0-2 px outward per edge, except the two "
         "'Box only, 0-10/0-20 px train jitter' rows.")
GEN = f"Generated {META['generated']} from the GAIVI exports, repository commit {META['commit']}."

wb = load_workbook(FULL)
ws1 = wb.create_sheet("Prostate (TRUS)", 0)
P = Page(ws1)
P.text("BAP-MOS camera-ready results: prostate, pooled TRUS", Font(bold=True, size=15))
for t in (GEN, COMMON, "Distances in mm (pixel spacing 0.159 mm). Test set: 39 slices from the simulation, "
          "case1 and case2 sites, pooled.", TRAIN):
    P.text(t, NOTE)
P.r += 1
P.block("1. Exact boxes (paper Table III uses the PTV columns)",
        ["Promptable models use exact, ground-truth-derived prompts; BAP-MOS uses its highest-Q action.",
         "SAM (box) = the Table I multi-organ / box / decoder-only cell (CE+Dice). Box only, BAP-MOS and the "
         "jitter variants use the BAP-MOS loss and protocol."],
        rows_models("TRUS", "exact", "exact",
                    ["nnunet", "unet", "sam_box", "medsam_box", "box_only", "reference", "medsam_reference",
                     "reference_evaltrain", "box_only_j10", "box_only_j20"]))
P.block("2. Ablations and matched controls (paper Table IV)",
        ["Decoder-only SAM, BAP-MOS loss and protocol, exact prompts."],
        rows_models("TRUS", "exact", "exact",
                    ["reference", "global", "dice_reward", "fixed_ring", "no_boundary_loss", "probe_disjoint",
                     "box_only", "fixed_235", "fixed_111", "fixed_511", "fixed_151"]))
P.block("3. Robustness to imperfect boxes (paper Table V and the robustness figure)",
        ["Identical perturbed boxes for every model (checked). Shifted ±n: every box edge (and point) moved by "
         "U{-n..n} px. Enlarged +n: every edge moved outward by U{0..n} px, points unchanged. 20 px ≈ 3.2 mm.",
         "The paper reports PTV and All-organs HD95 for the first five models (Table V); the remaining rows are the "
         "training-mixture tests of Table 'Where the robustness comes from' (fixed schedules, shared selector, "
         "box only with larger train jitter). A fixed schedule is tested with its highest-ratio prompt: 1:1:1 and "
         "5:1:1 with the box, 2:3:5 with box+point, 1:5:1 with the point (so 1:5:1 is unchanged by enlarged boxes)."],
        rows_conditions("TRUS", ["box_only", "sam_box", "reference", "medsam_box", "medsam_reference",
                                 "fixed_235", "global", "box_only_j10", "box_only_j20",
                                 "fixed_111", "fixed_511", "fixed_151"]))
rows = [["Setting", "Model"] + hdr_cols("TRUS")]
for m in ("box_only", "reference", "reference_evaltrain"):
    rows.append([L.COND["unet_boxes"], L.NAMES[m]] + [L.cell("TRUS", "auto_unet", "unet_boxes", m, o, me)
                                                      for o in L.ORGANS["TRUS"] for me in L.agg.METRICS])
for k in ("case1", "case2"):
    for m in ("box_only", "reference", "unet"):
        rows.append([L.COND[k], L.NAMES[m]] + [L.cell("TRUS", "loco", k, m, o, me)
                                               for o in L.ORGANS["TRUS"] for me in L.agg.METRICS])
P.block("4. Fully automatic prompts and leave-one-cohort-out (bottom of paper Table V)",
        ["U-Net boxes: prompts from the seed-matched U-Net prediction, pooled test set. LOCO: one clinical case "
         "held out of training and tested alone, exact prompts; the U-Net rows use no prompts."], rows)
P.block("5. Decoder-only vs full fine-tuning (paper Table I uses the PTV columns)",
        ["Fixed box or point prompts, CE+Dice, 2 px box jitter. The multi-organ / box / decoder-only cell is also "
         "the SAM (box) row. Single-organ cells segment the PTV only. Table II (prompt ratios) keeps its original "
         "numbers and is not in this workbook."],
        rows_models("TRUS", "table1_rerun", "exact",
                    ["single_points_decoder", "single_points_full", "single_box_decoder", "single_box_full",
                     "multi_points_decoder", "multi_points_full", "multi_box_decoder", "multi_box_full"]))
rows = [["Boxes", "Model"] + hdr_cols("TRUS")]
for c in ("noise5", "noise10", "noise20", "loose5", "loose10", "loose20"):
    for m in ("box_only", "reference", "reference_evaltrain"):
        if d["exports"].get(f"TRUS|superseded|{c}|{m}"):
            rows.append([L.COND[c], L.NAMES[m]] + [L.cell("TRUS", "superseded", c, m, o, me)
                                                   for o in L.ORGANS["TRUS"] for me in L.agg.METRICS])
P.block("6. First robustness version, unpaired (superseded by section 3)",
        ["Offsets from the same distribution but not identical boxes across models; the paper says it agrees "
         "with the paired test."], rows)
P.block("7. Paired tests",
        ["Two-sided Wilcoxon signed-rank on slice-organ pairs (each model's metric first averaged over its 3 "
         "seeds); pairs where either model has no valid boundary are dropped (A/B invalid). Holm across the "
         "organ tests of one comparison and metric; 'All organs' pools every pair (no Holm). Slices of one "
         "volume are correlated, so these p-values are optimistic. 'Fraction A better' = share of pairs "
         "with higher Dice / lower HD95 or MSD for A."], rows_paired("TRUS"), PFMT)
P.block("8. Checks before writing this sheet", [],
        [["Check", "Result"]] + [[k, v] for k, v in META["checks"]])

ws2 = wb.create_sheet("Bladder (PFUS1)", 1)
B = Page(ws2)
B.text("BAP-MOS camera-ready results: bladder, PFUS1", Font(bold=True, size=15))
for t in (GEN, COMMON, "Distances in pixels (pixel spacing 1.0). Patient-level split; 2101 test frames. "
          "Checkpoints are selected on the bladder.", TRAIN):
    B.text(t, NOTE)
B.r += 1
B.block("1. Exact boxes (the paper's PFUS1 table uses the Bladder columns)",
        ["U-Net uses no prompts. Each BAP-MOS variant is compared with the box-trained decoder of the same "
         "initialisation."],
        rows_models("PFUS1", "exact", "exact", ["unet", "sam_box", "bapmos_sam", "medsam_box", "bapmos_medsam"]))
B.block("2. Robustness to imperfect boxes",
        ["Same paired design as prostate: identical perturbed boxes for every model (checked). U-Net is "
         "constant across conditions (section 1)."],
        rows_conditions("PFUS1", ["sam_box", "bapmos_sam", "medsam_box", "bapmos_medsam"]))
B.block("3. Paired tests (patient level)",
        ["Two-sided Wilcoxon signed-rank on patients: each model's frame metrics are averaged over its 3 seeds, "
         "then over each patient's frames. Holm across the organ tests of one comparison and metric; 'All "
         "organs' is not Holm-corrected."], rows_paired("PFUS1"), PFMT)
B.block("4. Checks before writing this sheet", [],
        [["Check", "Result"]] + [[k, v] for k, v in META["checks"] if "PFUS1" in v or "Per-seed" in k
                                 or "SD definition" in k])

for ws in (ws1, ws2):
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 32
    for j in range(3, 30):
        ws.column_dimensions[get_column_letter(j)].width = 16
wb.active = 0  # open on the prostate sheet
wb.save(OUT)
print("saved", OUT)
