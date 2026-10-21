"""Build the camera-ready results workbook from collected.json (aggregate numbers only)."""
import datetime as dt
import json
import math
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import agg  # noqa: E402

OUT = Path(sys.argv[1])
META = json.loads(Path(sys.argv[2]).read_text())  # commit, generated, checks
d = agg.load()

NAMES = {
    "nnunet": "nnU-Net", "unet": "U-Net", "sam_box": "SAM (box)", "medsam_box": "MedSAM (box)",
    "box_only": "Box only (BAP-MOS loss)", "reference": "BAP-MOS (SAM)",
    "medsam_reference": "BAP-MOS (MedSAM)",
    "reference_evaltrain": "BAP-MOS (SAM), train-metrics pass on",
    "global": "Shared global selector", "dice_reward": "Dice-only reward",
    "fixed_ring": "Fixed 10 px ring (r_min)", "no_boundary_loss": "No boundary loss (alpha=1)",
    "probe_disjoint": "Disjoint probe / selection", "fixed_235": "Fixed 2:3:5 (same loss)",
    "fixed_111": "Fixed 1:1:1 (same loss; box at test)", "fixed_511": "Fixed 5:1:1 (same loss; box at test)",
    "fixed_151": "Fixed 1:5:1 (same loss; point at test)",
    "box_only_j10": "Box only, 0-10 px train jitter", "box_only_j20": "Box only, 0-20 px train jitter",
    "bapmos_sam": "BAP-MOS (SAM)", "bapmos_medsam": "BAP-MOS (MedSAM)",
    "single_points_decoder": "Single-organ, points, decoder-only",
    "single_points_full": "Single-organ, points, full fine-tuning",
    "single_box_decoder": "Single-organ, box, decoder-only",
    "single_box_full": "Single-organ, box, full fine-tuning",
    "multi_points_decoder": "Multi-organ, points, decoder-only",
    "multi_points_full": "Multi-organ, points, full fine-tuning",
    "multi_box_decoder": "Multi-organ, box, decoder-only",
    "multi_box_full": "Multi-organ, box, full fine-tuning",
}
COND = {
    "exact": "Exact boxes", "unet_boxes": "Boxes from U-Net masks",
    "case1": "LOCO: clinical case 1 held out", "case2": "LOCO: clinical case 2 held out",
    **{f"pshift{n}": f"Shifted ±{n} px" for n in (5, 10, 20)},
    **{f"ploose{n}": f"Enlarged +{n} px" for n in (5, 10, 20)},
    **{f"noise{n}": f"Shifted ±{n} px, unpaired (superseded)" for n in (5, 10, 20)},
    **{f"loose{n}": f"Enlarged +{n} px, unpaired (superseded)" for n in (5, 10, 20)},
}
ROBUST = ["exact"] + [f"pshift{n}" for n in (5, 10, 20)] + [f"ploose{n}" for n in (5, 10, 20)]
ORGANS = {"TRUS": ["PTV", "Bladder", "Rectum", "Urethra", "ALL"],
          "PFUS1": ["Bladder", "Anus", "Levator ani muscle", "Pubis", "Rectum", "Urethra", "Uterus",
                    "Vagina", "ALL"]}
UNIT = {"TRUS": "mm", "PFUS1": "px"}
MLABEL = {"dice": "Dice", "hd95_mm": "HD95", "msd_mm": "MSD"}


def decimals(ds, metric):
    return 3 if (metric == "dice" or ds == "TRUS") else 2


def fmt(a, ds, metric):
    if a is None or a[3] < 3:
        return None
    k = decimals(ds, metric)
    return f"{a[0]:.{k}f} ± {a[1]:.{k}f}"


def entry(ds, block, cond, model):
    return d["exports"].get(f"{ds}|{block}|{cond}|{model}", {})


def cell(ds, block, cond, model, organ, metric):
    e = entry(ds, block, cond, model)
    a = agg.agg(e, organ, metric) if e else None
    return fmt(a, ds, metric) or ("pending" if not e or len(e) < 3 else "n/a")


HDR = Font(bold=True)
HFILL = PatternFill("solid", fgColor="DDE5F0")
PEND = Font(color="9C0006", italic=True)
wb = Workbook()


def sheet(title, rows, widths=None, notes=()):
    ws = wb.create_sheet(title)
    r0 = 1
    for n in notes:
        ws.cell(row=r0, column=1, value=n).font = Font(italic=True)
        r0 += 1
    if notes:
        r0 += 1
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = ws.cell(row=r0 + i, column=j + 1, value=v)
            if i == 0:
                c.font, c.fill = HDR, HFILL
                c.alignment = Alignment(wrap_text=True, vertical="top")
            elif v == "pending":
                c.font = PEND
    ws.freeze_panes = ws.cell(row=r0 + 1, column=3 if len(rows[0]) > 2 else 2)
    for j in range(len(rows[0])):
        w = (widths[j] if widths and j < len(widths) else
             min(48, max(10, max(len(str(r[j])) for r in rows if j < len(r) and r[j] is not None) + 2)))
        ws.column_dimensions[get_column_letter(j + 1)].width = w
    return ws


def table(ds, block, cond, models, organs, title, notes):
    hdr = ["Model"] + [f"{o if o != 'ALL' else 'All organs'} {MLABEL[m]}" + ("" if m == "dice" else f" ({UNIT[ds]})")
                       for o in organs for m in agg.METRICS]
    rows = [hdr]
    for m in models:
        rows.append([NAMES[m]] + [cell(ds, block, cond, m, o, me) for o in organs for me in agg.METRICS])
    return sheet(title, rows, [34] + [17] * (len(hdr) - 1), notes)


COMMON = ("Cells: mean ± sample SD (ddof=1) over seeds 42/43/44; 'pending' = fewer than 3 seeds "
          "finished. Per seed: mean over slice-organ pairs with a valid boundary, pooled across sites.")

# README is filled last (needs check results); create it first so it is the first tab.
readme = wb.active
readme.title = "README"

table("TRUS", "exact", "exact",
      ["nnunet", "unet", "sam_box", "medsam_box", "box_only", "reference", "medsam_reference",
       "reference_evaltrain", "box_only_j10", "box_only_j20"],
      ORGANS["TRUS"], "TRUS exact boxes",
      ["Pooled TRUS test set (39 slices; sites simulation/case1/case2). Promptable models use exact, "
       "ground-truth-derived prompts (BAP-MOS: highest-Q action). Paper Table III uses the PTV columns.",
       "SAM (box) = Table I multi-organ / box / decoder-only cell (CE+Dice). Box only, BAP-MOS (SAM), "
       "BAP-MOS (MedSAM) and the jitter variants use the BAP-MOS loss and protocol.", COMMON])

table("TRUS", "exact", "exact",
      ["reference", "global", "dice_reward", "fixed_ring", "no_boundary_loss", "probe_disjoint",
       "box_only", "fixed_235", "fixed_111", "fixed_511", "fixed_151"],
      ORGANS["TRUS"], "TRUS ablations",
      ["Paper Table IV (frozen in the paper; these are the camera-ready runs it was built from). "
       "Decoder-only SAM, BAP-MOS loss and protocol, exact prompts.", COMMON])


def robust_sheet(ds, models, title, notes):
    hdr = ["Boxes", "Model"]
    organs = ORGANS[ds]
    hdr += [f"{o if o != 'ALL' else 'All organs'} {MLABEL[m]}" + ("" if m == "dice" else f" ({UNIT[ds]})")
            for o in organs for m in agg.METRICS]
    rows = [hdr]
    for c in ROBUST:
        for m in models:
            blk = "exact" if c == "exact" else "robust"
            rows.append([COND[c], NAMES[m]] + [cell(ds, blk, c, m, o, me) for o in organs for me in agg.METRICS])
    return sheet(title, rows, [20, 32] + [17] * (len(hdr) - 2), notes)


robust_sheet("TRUS", ["box_only", "sam_box", "reference", "medsam_box", "medsam_reference", "fixed_235",
                      "global", "box_only_j10", "box_only_j20", "fixed_111", "fixed_511", "fixed_151"], "TRUS robustness",
             ["Paired test-time perturbations: identical perturbed boxes for every model (checked). "
              "Shifted ±n: every box edge (and point) moved by U{-n..n} px. Enlarged +n: every edge moved "
              "outward by U{0..n} px, points unchanged. 20 px ≈ 3.2 mm. Exact = the exact-box exports.",
              "Paper Table V / Fig. robustness use PTV HD95 and All-organs HD95 for the first five models.",
              COMMON])

# LOCO and automatic prompts
hdr = ["Setting", "Model"] + [f"{o if o != 'ALL' else 'All organs'} {MLABEL[m]}" + ("" if m == "dice" else " (mm)")
                              for o in ORGANS["TRUS"] for m in agg.METRICS]
rows = [hdr]
for m in ("box_only", "reference", "reference_evaltrain"):
    rows.append([COND["unet_boxes"], NAMES[m]] + [cell("TRUS", "auto_unet", "unet_boxes", m, o, me)
                                                  for o in ORGANS["TRUS"] for me in agg.METRICS])
for k in ("case1", "case2"):
    for m in ("box_only", "reference", "unet"):
        rows.append([COND[k], NAMES[m]] + [cell("TRUS", "loco", k, m, o, me)
                                           for o in ORGANS["TRUS"] for me in agg.METRICS])
sheet("TRUS LOCO + U-Net boxes", rows, [34, 30] + [17] * (len(hdr) - 2),
      ["U-Net boxes: prompts derived from the seed-matched U-Net prediction (fully automatic pipeline); "
       "pooled test set. LOCO: one clinical case held out of training, tested on that case only, exact "
       "prompts; U-Net rows are the LOCO U-Net without prompts.", COMMON])

table("TRUS", "table1_rerun", "exact",
      ["single_points_decoder", "single_points_full", "single_box_decoder", "single_box_full",
       "multi_points_decoder", "multi_points_full", "multi_box_decoder", "multi_box_full"],
      ORGANS["TRUS"], "TRUS Table I rerun",
      ["Paper Table I (tab:decoder_vs_full): decoder-only vs full fine-tuning with fixed box or point prompts "
       "(CE+Dice, 2 px box jitter), camera-ready runs; the paper reports the PTV columns. The multi-organ / box / "
       "decoder-only cell is also the SAM (box) row. Table II (prompt ratios) keeps its original numbers and is not here.",
       "Single-organ cells segment the PTV only (other organs n/a; their All-organs column equals the PTV).", COMMON])

table("PFUS1", "exact", "exact", ["unet", "sam_box", "bapmos_sam", "medsam_box", "bapmos_medsam"],
      ORGANS["PFUS1"], "PFUS1 exact boxes",
      ["PFUS1, patient-level split (2101 test frames). Distances in pixels (pixel spacing 1.0). "
       "Checkpoints selected on the bladder; paper Table (tab:pfus1) uses the Bladder columns.", COMMON])

robust_sheet("PFUS1", ["sam_box", "bapmos_sam", "medsam_box", "bapmos_medsam"], "PFUS1 robustness",
             ["Same paired perturbation design as TRUS (identical boxes for every model, checked). "
              "Distances in pixels. U-Net uses no prompts and is constant across conditions "
              "(see PFUS1 exact boxes).", COMMON])

hdr = ["Boxes", "Model"] + [f"{o if o != 'ALL' else 'All organs'} {MLABEL[m]}" + ("" if m == "dice" else " (mm)")
                            for o in ORGANS["TRUS"] for m in agg.METRICS]
rows = [hdr]
for c in ("noise5", "noise10", "noise20", "loose5", "loose10", "loose20"):
    for m in ("box_only", "reference", "reference_evaltrain"):
        if f"TRUS|superseded|{c}|{m}" in d["exports"] and d["exports"][f"TRUS|superseded|{c}|{m}"]:
            rows.append([COND[c], NAMES[m]] + [cell("TRUS", "superseded", c, m, o, me)
                                               for o in ORGANS["TRUS"] for me in agg.METRICS])
sheet("TRUS superseded", rows, [36, 34] + [17] * (len(hdr) - 2),
      ["First robustness version: offsets from the same distribution but NOT identical boxes across models. "
       "Superseded by the paired test (TRUS robustness); kept because the paper mentions it agrees.", COMMON])

# Paired tests
hdr = ["Dataset", "Boxes", "A", "B", "Unit", "Organ", "Metric", "n pairs", "A mean", "B mean",
       "Median diff (A-B)", "Fraction A better", "p (Wilcoxon)", "p (Holm, across organs)",
       "A invalid", "B invalid", "Status"]
rows = [hdr]
for name, p in d["paired"].items():
    ds, cond, pair = name.split("|")
    a, b = p["a"].split("|")[3], p["b"].split("|")[3]
    base = [ds, COND.get(cond.replace("loco_", ""), cond), NAMES[a], NAMES[b]]
    if p["status"] != "ok":
        rows.append(base + [("patient" if ds == "PFUS1" else "slice")] + [None] * 11 + ["pending"])
        continue
    for r in p["rows"]:
        nan = lambda x: None if x is None or (isinstance(x, float) and math.isnan(x)) else x
        rows.append(base + [p["unit"], "All organs" if r["organ"] == "ALL" else r["organ"], MLABEL[r["metric"]],
                            r["n_pairs"], nan(r["A_mean"]), nan(r["B_mean"]), nan(r["median_diff_a_minus_b"]),
                            nan(r["a_better_frac"]), nan(r["p_wilcoxon"]), nan(r["p_holm"]),
                            r["a_invalid"], r["b_invalid"], "ok"])
ws = sheet("Paired tests", rows, [8, 30, 30, 30, 8, 18, 7, 8, 10, 10, 12, 10, 12, 14, 8, 8, 9],
           ["Two-sided Wilcoxon signed-rank (bapmos.results.paired_stats, unchanged). Each slice-organ "
            "metric is first averaged over the 3 seeds of each model; pairs where either model has no valid "
            "boundary are dropped (counts in 'A/B invalid'). Holm correction across the organ tests of one "
            "comparison and metric; 'All organs' pools every slice-organ pair and is not Holm-corrected.",
            "PFUS1: unit = patient (frames averaged per patient first). TRUS: unit = slice; slices of one "
            "volume are correlated, so TRUS p-values are optimistic (stated in the paper).",
            "'Fraction A better' = share of pairs where A has higher Dice / lower HD95 or MSD."])
for row in ws.iter_rows(min_row=next(c.row for c in ws["A"] if c.value == "Dataset") + 1):
    for c in row[8:12]:
        c.number_format = "0.0000"
    for c in row[12:14]:
        c.number_format = "0.00E+00"

# Long format: everything, numeric
hdr = ["Dataset", "Block", "Boxes", "Model", "Organ", "Metric", "Unit", "Mean", "SD", "Mean ± SD",
       "Seed 42", "Seed 43", "Seed 44", "Seeds finished", "Valid pairs (seed 42)", "Export (seed 42)"]
rows = [hdr]
for key, e in d["exports"].items():
    ds, block, cond, model = key.split("|")
    if block == "superseded" and not e:
        continue  # never planned (unpaired enlarged boxes were not run for the train-metrics variant)
    for o in ORGANS[ds]:
        for m in agg.METRICS:
            a = agg.agg(e, o, m) if e else None
            per = {s: (e[s]["metrics"][o][m] if s in e and o in e[s]["metrics"] else None) for s in agg.SEEDS}
            rows.append([ds, block, COND.get(cond, cond), NAMES[model], "All organs" if o == "ALL" else o,
                         MLABEL[m], "" if m == "dice" else UNIT[ds],
                         a[0] if a and a[3] == 3 else None, a[1] if a and a[3] == 3 else None,
                         fmt(a, ds, m) or ("n/a" if len(e) == 3 else "pending"),
                         per["seed42"], per["seed43_rep2"], per["seed44_rep3"], len(e),
                         e["seed42"]["metrics"][o]["n"] if "seed42" in e and o in e["seed42"]["metrics"] else None,
                         e["seed42"]["path"].split("camera_ready/")[1] if "seed42" in e else None])
ws = sheet("All results (long)", rows, [8, 12, 30, 34, 18, 7, 5, 10, 10, 18, 10, 10, 10, 8, 10, 48],
           ["One row per dataset × condition × model × organ × metric. Numeric columns for filtering and "
            "pivoting; Mean and SD are left empty unless all 3 seeds finished.", COMMON])
hr = next(c.row for c in ws["A"] if c.value == "Dataset")
for row in ws.iter_rows(min_row=hr + 1):
    for c in (row[7], row[8], row[10], row[11], row[12]):
        c.number_format = "0.0000"
ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(hdr))}{ws.max_row}"

# Checks and pending
rows = [["Check", "Result"]] + [[k, v] for k, v in META["checks"]]
sheet("Checks", rows, [70, 70])
pend = sorted(p.split("camera_ready/")[1] for p in d["missing"]
              if not re.search(r"/loose\d+/reference_evaltrain_", p))  # never planned
sheet("Pending exports", [["Export folder (seed)"]] + [[p] for p in pend], [70],
      [f"{len(pend)} seed exports not finished when this workbook was generated."])

# README
lines = [
    ("BAP-MOS camera-ready results workbook", True),
    (f"Generated {META['generated']} from GAIVI exports; repository commit {META['commit']} "
     "(branch camera-ready-experiments).", False),
    ("", False),
    ("Every cell is mean ± sample SD (ddof=1) over seeds 42/43/44. A cell reads 'pending' until all three "
     "seeds have finished; partial results are never averaged.", False),
    ("Per seed: the metric is averaged over the slice-organ pairs with a valid boundary, pooled across test "
     "sites (TRUS: simulation, case1, case2). 'All organs' averages every such pair.", False),
    ("Units: TRUS distances in mm (pixel spacing 0.159 mm); PFUS1 distances in pixels (spacing 1.0).", False),
    ("Only aggregate metrics are included: no per-slice rows, no image identifiers.", False),
    ("", False),
    ("Sheets", True),
    ("TRUS exact boxes: Table III and the per-organ numbers quoted in the text.", False),
    ("TRUS ablations: Table IV (frozen in the paper).", False),
    ("TRUS robustness: Table V and Fig. robustness, plus the attribution rows (fixed 2:3:5, shared selector, "
     "box only with larger train jitter).", False),
    ("TRUS LOCO + U-Net boxes: the bottom block of Table V.", False),
    ("TRUS Table I rerun: paper Table I (decoder-only vs full fine-tuning); Table II (prompt ratios) keeps its original numbers.", False),
    ("PFUS1 exact boxes / PFUS1 robustness: the PFUS1 table and robustness results.", False),
    ("TRUS superseded: the first, unpaired shifted/enlarged-box exports.", False),
    ("Paired tests: every Wilcoxon comparison, per organ and over all organs.", False),
    ("All results (long): every number in numeric form with per-seed values (filterable).", False),
    ("Checks: the verification run before this workbook was written. Pending exports: what is still running.", False),
    ("", False),
    ("Train-time box jitter: 0-2 px outward per edge for every box-prompted model on both datasets, except "
     "the two 'Box only, 0-10/0-20 px train jitter' rows.", False),
]
for i, (t, bold) in enumerate(lines, 1):
    c = readme.cell(row=i, column=1, value=t)
    c.font = Font(bold=bold, size=13 if i == 1 else 11)
readme.column_dimensions["A"].width = 130
wb.save(OUT)
print("saved", OUT)
