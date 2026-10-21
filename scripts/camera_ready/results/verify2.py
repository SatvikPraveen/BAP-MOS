"""Check the two dataset sheets cell by cell against the long sheet and the collected paired tests."""
import math, re, sys
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0, str(Path(__file__).resolve().parent))
import labels as L
wb = load_workbook(sys.argv[1]); d = L.d
long = wb["All results (long)"]
HR = next(c.row for c in long["A"] if c.value == "Dataset")
h = [c.value for c in long[HR]]; ix = {k: i for i, k in enumerate(h)}
lk = {}
for r in long.iter_rows(min_row=HR + 1, values_only=True):
    lk.setdefault((r[ix["Dataset"]], r[ix["Block"]], r[ix["Boxes"]], r[ix["Model"]], r[ix["Organ"]], r[ix["Metric"]]), r[ix["Mean ± SD"]])
problems, nchk, npair = [], 0, 0
SECT_BLOCK = {"1. Exact boxes": "exact", "2. Ablations": "exact", "3. Robustness": None, "4. Fully automatic": None,
              "5. Decoder-only": "table1_rerun", "6. First robustness": "superseded", "2. Robustness": None}
BLK_OF = {"Exact boxes": "exact", "Boxes from U-Net masks": "auto_unet"}
for title, ds in (("Prostate (TRUS)", "TRUS"), ("Bladder (PFUS1)", "PFUS1")):
    ws = wb[title]; rows = list(ws.iter_rows(values_only=True)); sect = None; hdr = None; paired_rows = []
    for row in rows:
        v0 = row[0]
        if isinstance(v0, str) and re.match(r"^\d\. ", v0):
            sect, hdr = v0, None; continue
        if v0 in ("Model", "Boxes", "Setting", "Check"):
            hdr = list(row); continue
        if hdr is None or v0 is None or sect is None:
            continue
        for x in row:
            if isinstance(x, str) and re.search(r"\b(nan|None|inf)\b", x): problems.append(f"{title}: text {x}")
            if isinstance(x, float) and not math.isfinite(x): problems.append(f"{title}: non-finite")
        if "Paired tests" in sect:
            paired_rows.append(row); continue
        if hdr[0] == "Check": continue
        key = next((k for k in SECT_BLOCK if sect.startswith(k)), None)
        if key is None: problems.append(f"unknown section {sect}"); continue
        boxes, model = (row[0], row[1]) if hdr[0] in ("Boxes", "Setting") else ("Exact boxes", row[0])
        blk = SECT_BLOCK[key]
        if blk is None:
            blk = BLK_OF.get(boxes, "loco" if boxes.startswith("LOCO") else "robust")
        for j in range(2, len(hdr)):
            m = re.match(r"^(.*) (Dice|HD95|MSD)", hdr[j])
            want = lk.get((ds, blk, boxes, model, m.group(1), m.group(2)))
            nchk += 1
            if want is None or want != row[j]:
                problems.append(f"{title} {sect[:12]} {boxes}/{model}/{m.group(1)}/{m.group(2)}: {row[j]} vs {want}")
    # paired: compare with collected rows in order
    src = []
    for name, p in d["paired"].items():
        if name.split("|")[0] != ds: continue
        src += p["rows"] if p["status"] == "ok" else [None]
    if len(src) != len(paired_rows): problems.append(f"{title}: paired rows {len(paired_rows)} vs {len(src)}")
    for r, s in zip(paired_rows, src):
        npair += 1
        if s is None:
            if r[14] != "pending": problems.append(f"{title}: pending not marked")
            continue
        for col, f in ((5, "n_pairs"), (9, "a_better_frac"), (10, "p_wilcoxon"), (11, "p_holm")):
            a, b = r[col], s[f]
            b = None if isinstance(b, float) and math.isnan(b) else b
            if (a is None) != (b is None) or (a is not None and abs(a - b) > 1e-15):
                problems.append(f"{title}: paired {r[:5]} {f} {a} vs {b}")
print(f"dataset-sheet cells checked: {nchk}; paired rows checked: {npair}; PROBLEMS: {len(problems)}")
for p in problems[:20]: print(" ", p)
